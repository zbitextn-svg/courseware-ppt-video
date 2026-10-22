# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
p = Presentation(path)


def walk(shapes, depth=0):
    out = []
    for sh in shapes:
        pad = "  " * depth
        if sh.shape_type == 6 or sh.__class__.__name__ == "GroupShape":
            out.append(f"{pad}[GROUP {sh.name}]")
            try:
                out += walk(sh.shapes, depth + 1)
            except Exception:
                pass
            continue
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    lv = para.level
                    sz = None
                    for r in para.runs:
                        if r.font.size:
                            sz = r.font.size.pt
                            break
                    out.append(f"{pad}TXT lv{lv} {sz}pt | {t}")
        if getattr(sh, "has_table", False):
            tb = sh.table
            out.append(f"{pad}[TABLE {len(tb.rows)}x{len(tb.columns)}]")
            for row in tb.rows:
                cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
                out.append(f"{pad}  | " + " | ".join(cells))
        if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
            try:
                w = Emu(sh.width).inches
                h = Emu(sh.height).inches
            except Exception:
                w = h = 0
            out.append(f"{pad}[PIC {sh.name} {w:.2f}x{h:.2f}in]")
    return out


for i, s in enumerate(p.slides, 1):
    print(f"\n===== SLIDE {i} =====")
    for line in walk(s.shapes):
        print(line)
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        print("  NOTES:", s.notes_slide.notes_text_frame.text.strip()[:300])
